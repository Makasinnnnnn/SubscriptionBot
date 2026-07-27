from collections.abc import Iterable, Sequence
from datetime import UTC, datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

MAX_EXPORT_ROWS = 5000
FORMULA_PREFIXES = ("=", "+", "-", "@")


def safe_excel_value(value: Any) -> Any:
    if isinstance(value, datetime):
        aware = value.replace(tzinfo=UTC) if value.tzinfo is None else value.astimezone(UTC)
        return aware.replace(tzinfo=None)
    if isinstance(value, str) and value.startswith(FORMULA_PREFIXES):
        return f"'{value}"
    return value


def build_xlsx(
    *,
    sheet_name: str,
    headers: Sequence[str],
    rows: Iterable[Sequence[Any]],
) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = _safe_sheet_name(sheet_name)
    sheet.freeze_panes = "A2"
    sheet.append(list(headers))

    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")

    row_count = 0
    for row in rows:
        if row_count >= MAX_EXPORT_ROWS:
            break
        sheet.append([safe_excel_value(value) for value in row])
        row_count += 1
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{max(1, sheet.max_row)}"
    )

    for column_index, header in enumerate(headers, start=1):
        values = [
            str(sheet.cell(row=row, column=column_index).value or "")
            for row in range(1, min(sheet.max_row, 200) + 1)
        ]
        width = min(45, max(10, len(str(header)) + 2, *(len(value) + 2 for value in values)))
        sheet.column_dimensions[get_column_letter(column_index)].width = width

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, datetime):
                cell.number_format = "yyyy-mm-dd hh:mm:ss"
            cell.alignment = Alignment(vertical="top")

    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


def _safe_sheet_name(value: str) -> str:
    cleaned = "".join("_" if character in r"[]:*?/\\" else character for character in value)
    return cleaned[:31] or "Export"
