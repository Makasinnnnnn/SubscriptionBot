from datetime import datetime
from io import BytesIO

from openpyxl import load_workbook

from app.services.xlsx_export_service import build_xlsx


def test_xlsx_export_has_filters_dates_and_formula_protection() -> None:
    payload = build_xlsx(
        sheet_name="Платежи/2026",
        headers=["ID", "Username", "Создан"],
        rows=[
            (1, "=HYPERLINK(\"bad\")", datetime(2026, 7, 27, 12, 30)),
            (2, "+cmd", datetime(2026, 7, 27, 13, 30)),
            (3, "-formula", datetime(2026, 7, 27, 14, 30)),
            (4, "@value", datetime(2026, 7, 27, 15, 30)),
        ],
    )

    workbook = load_workbook(BytesIO(payload))
    sheet = workbook.active

    assert sheet.title == "Платежи_2026"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref == "A1:C5"
    assert [cell.value for cell in sheet[1]] == ["ID", "Username", "Создан"]
    assert [sheet.cell(row=row, column=2).value for row in range(2, 6)] == [
        "'=HYPERLINK(\"bad\")",
        "'+cmd",
        "'-formula",
        "'@value",
    ]
    assert sheet["C2"].number_format == "yyyy-mm-dd hh:mm:ss"
    workbook.close()
