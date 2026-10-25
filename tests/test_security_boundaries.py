import asyncio
import hashlib
import hmac
import io
from types import SimpleNamespace
from unittest.mock import patch

import pytest
from fastapi import HTTPException
from fastapi.security import HTTPBasicCredentials
from openpyxl import load_workbook
from starlette.requests import Request

from app.admin import routes as admin_routes
from app.api.routes import webhooks
from app.core import rate_limit
from app.payments.providers import CryptoBotProvider, WebhookVerificationError
from app.services.xlsx_export_service import build_xlsx


def request(body=b"{}", headers=(), chunks=None):
    chunks = list(chunks or [body])

    async def receive():
        value = chunks.pop(0)
        return {"type": "http.request", "body": value, "more_body": bool(chunks)}

    return Request(
        {
            "type": "http",
            "method": "POST",
            "path": "/",
            "client": ("192.0.2.1", 40000),
            "headers": list(headers),
        },
        receive,
    ), chunks


def test_auth_rate_limit_cannot_be_bypassed_with_username_rotation():
    async def scenario():
        limiter = rate_limit.SlidingWindowLimiter()
        settings = SimpleNamespace(admin_password="test-password-123456", admin_username="admin")
        codes = []
        with (
            patch.object(admin_routes, "get_settings", return_value=settings),
            patch.object(admin_routes, "auth_limiter", limiter),
        ):
            for n in range(20):
                req, _ = request()
                try:
                    await admin_routes.require_web_admin(
                        req, HTTPBasicCredentials(username=f"wrong-{n}", password="bad")
                    )
                except HTTPException as error:
                    codes.append(error.status_code)
        assert codes[10:] == [429] * 10, (
            f"username rotation results={codes} buckets={len(limiter._events)}"
        )

    asyncio.run(scenario())


def test_limiter_reclaims_expired_inactive_keys():
    async def scenario():
        limiter = rate_limit.SlidingWindowLimiter()
        with patch.object(rate_limit, "monotonic", return_value=1.0):
            for n in range(1000):
                await limiter.allow(str(n), limit=10, window_seconds=1)
        with patch.object(rate_limit, "monotonic", return_value=100.0):
            await limiter.allow("new-key", limit=10, window_seconds=1)
        assert len(limiter._events) <= 1, f"expired buckets retained={len(limiter._events)}"

    asyncio.run(scenario())


def test_limiter_refuses_new_keys_when_capacity_is_reached():
    async def scenario():
        limiter = rate_limit.SlidingWindowLimiter(max_keys=2)
        assert await limiter.allow("one", limit=10, window_seconds=60)
        assert await limiter.allow("two", limit=10, window_seconds=60)
        assert not await limiter.allow("three", limit=10, window_seconds=60)
        assert await limiter.is_blocked("three", limit=10, window_seconds=60)
        assert len(limiter._events) == 2
        assert await limiter.allow("one", limit=10, window_seconds=60)

    asyncio.run(scenario())


def test_stream_body_limit_stops_before_reading_all_chunks():
    async def scenario():
        req, remaining = request(
            headers=[(b"content-type", b"application/json")], chunks=[b"x" * 8 for _ in range(50)]
        )
        settings = SimpleNamespace(
            webhook_secret="audit", yookassa_allowed_ips=[], webhook_max_body_bytes=16
        )
        with (
            patch.object(webhooks, "get_settings", return_value=settings),
            patch.object(webhooks, "webhook_limiter", rate_limit.SlidingWindowLimiter()),
            pytest.raises(HTTPException) as raised,
        ):
            await webhooks._process("cryptobot", req, None, "audit")
        assert raised.value.status_code == 413
        assert remaining, "body limit exhausted every chunk before rejecting"

    asyncio.run(scenario())


def test_invalid_utf8_json_returns_client_error():
    async def scenario():
        req, _ = request(body=b"\xff", headers=[(b"content-type", b"application/json")])
        settings = SimpleNamespace(
            webhook_secret="audit", yookassa_allowed_ips=[], webhook_max_body_bytes=16
        )
        with (
            patch.object(webhooks, "get_settings", return_value=settings),
            patch.object(webhooks, "webhook_limiter", rate_limit.SlidingWindowLimiter()),
            pytest.raises(HTTPException) as raised,
        ):
            await webhooks._process("cryptobot", req, None, "audit")
        assert raised.value.status_code == 400

    asyncio.run(scenario())


def test_xlsx_formula_input_is_never_written_as_formula():
    payloads = ["=1+1", "+1+1", "-1+1", "@SUM(A1)", "\t=1+1", " =1+1", "\n=1+1"]
    raw = build_xlsx(sheet_name="Audit", headers=("name",), rows=[(value,) for value in payloads])
    sheet = load_workbook(io.BytesIO(raw)).active
    assert all(sheet.cell(row=n, column=1).data_type != "f" for n in range(2, 2 + len(payloads)))


def test_crypto_signature_is_over_exact_original_bytes():
    token = "audit-token-no-provider-access"
    provider = CryptoBotProvider(token=token, api_url="https://synthetic.invalid")
    original = b'{ "test": 1 }'
    key = hashlib.sha256(token.encode()).digest()
    signature = hmac.new(key, original, hashlib.sha256).hexdigest()
    provider._verify_signature({"crypto-pay-api-signature": signature}, original)
    with pytest.raises(WebhookVerificationError):
        provider._verify_signature({"crypto-pay-api-signature": signature}, b'{"test":1}')


def test_proxy_headers_ignored_from_untrusted_direct_peer():
    req, _ = request(
        headers=[(b"x-forwarded-for", b"185.71.76.1"), (b"cf-connecting-ip", b"185.71.76.1")]
    )
    with patch.object(
        webhooks, "get_settings", return_value=SimpleNamespace(trusted_proxy_ips=["127.0.0.1/32"])
    ):
        assert webhooks._client_ip(req) == "192.0.2.1"


def test_scoped_admin_cannot_disable_global_materials_channel():
    from unittest.mock import AsyncMock

    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

    from app.admin.service import AdminPermissionService
    from app.admin.settings_service import RuntimeSettingsService
    from app.bot.handlers import channels
    from app.db.base import Base
    from app.db.models import Admin

    async def scenario():
        engine = create_async_engine("sqlite+aiosqlite:///:memory:")
        factory = async_sessionmaker(engine, expire_on_commit=False)
        async with engine.begin() as connection:
            await connection.run_sync(Base.metadata.create_all)
        async with factory() as session:
            session.add(Admin(telegram_id=700001, role="admin", all_channels=False))
            await RuntimeSettingsService(session).set(
                RuntimeSettingsService.MATERIALS_CHANNEL_ID, -100700001, 1
            )
            await session.commit()
            assert await AdminPermissionService(session).channel_scope(700001) == []
            assert not await AdminPermissionService(session).can_manage_all_channels(700001)
        callback = SimpleNamespace(
            from_user=SimpleNamespace(id=700001), data="chan:materialset:0", answer=AsyncMock()
        )
        with (
            patch.object(channels, "async_session_factory", factory),
            patch.object(channels, "_show_channels", AsyncMock()),
        ):
            await channels.materials_set(callback)
        async with factory() as session:
            remaining = await RuntimeSettingsService(session).get(
                RuntimeSettingsService.MATERIALS_CHANNEL_ID, None
            )
        await engine.dispose()
        assert remaining == -100700001, (
            f"scoped admin replaced global materials setting with {remaining!r}"
        )

    asyncio.run(scenario())


@pytest.mark.asyncio
async def test_payment_report_and_callback_scope_survive_plan_channel_change(session):
    from contextlib import asynccontextmanager
    from decimal import Decimal
    from unittest.mock import AsyncMock

    from app.bot.handlers import admin_reports
    from app.db.enums import PaymentStatus
    from app.db.models import Admin, AdminChannelAccess, Channel, Payment, Plan, User

    primary = Channel(telegram_chat_id=-100900001, title="Purchased channel")
    forbidden = Channel(telegram_chat_id=-100900002, title="Different channel")
    user = User(telegram_id=900001)
    administrator = Admin(telegram_id=900002, role="admin", all_channels=False)
    session.add_all([primary, forbidden, user, administrator])
    await session.flush()
    plan = Plan(
        channel_id=primary.id, title="Period", duration_days=30, price=Decimal(10), currency="RUB"
    )
    session.add(plan)
    await session.flush()
    payment = Payment(
        user_id=user.id,
        plan_id=plan.id,
        provider="dev",
        amount=Decimal(10),
        currency="RUB",
        idempotency_key="scope-test",
        purchased_channel_id=primary.id,
        purchased_duration_days=30,
        status=PaymentStatus.PAID,
    )
    session.add(payment)
    session.add(AdminChannelAccess(admin_id=administrator.id, channel_id=primary.id))
    await session.commit()
    plan.channel_id = forbidden.id
    await session.commit()
    session.expunge_all()
    visible = list(
        await session.scalars(admin_reports._payments_statement("all", "all", [primary.id]))
    )
    hidden = list(
        await session.scalars(admin_reports._payments_statement("all", "all", [forbidden.id]))
    )
    assert [item.id for item in visible] == [payment.id]
    assert hidden == []
    assert visible[0].access_channel_id == primary.id
    session.expunge_all()

    @asynccontextmanager
    async def factory():
        yield session

    callback = SimpleNamespace(
        from_user=SimpleNamespace(id=900002), data=f"rpt:payment:{payment.id}", answer=AsyncMock()
    )
    editor = AsyncMock()
    with (
        patch.object(admin_reports, "async_session_factory", factory),
        patch.object(admin_reports, "_edit", editor),
    ):
        await admin_reports.payment_card(callback)
    assert editor.await_count == 1
    assert "Purchased channel" in editor.await_args.args[1]
    assert "Different channel" not in editor.await_args.args[1]
